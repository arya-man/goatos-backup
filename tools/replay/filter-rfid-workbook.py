#!/usr/bin/env python3
"""Filter a live RFID workbook to a deterministic seed set.

The live RFID sheet may contain rows that are new tag assignments or review
deltas, not necessarily new goat passports. Until RFID apply can attach those
rows to existing old-tag passports, the replay seed must only auto-create rows
whose RFID existed in the trusted baseline export. New live rows remain visible
in diagnostics, but they do not silently inflate the herd.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook


NON_ALNUM = re.compile(r"[^A-Z0-9]+")


def canonical_identifier(raw: object) -> str:
    value = str(raw or "").strip().replace(",", "").replace(" ", "")
    if not value or value.lower() == "none":
        return ""
    if "e" in value.lower():
        try:
            value = str(int(float(value)))
        except ValueError:
            pass
    elif "." in value:
        left, _, right = value.partition(".")
        if right.strip("0") == "":
            value = left
    return NON_ALNUM.sub("", value.upper())


def read_baseline_rfids(path: Path) -> set[str]:
    rfids: set[str] = set()
    with path.open(newline="") as f:
        for row in csv.DictReader(f):
            value = canonical_identifier(row.get("RFID") or row.get("rfid"))
            if value:
                rfids.add(value)
    return rfids


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-xlsx", required=True, type=Path)
    parser.add_argument("--sheet", required=True)
    parser.add_argument("--baseline-csv", required=True, type=Path)
    parser.add_argument("--output-xlsx", required=True, type=Path)
    parser.add_argument("--excluded-csv", required=True, type=Path)
    parser.add_argument("--summary-json", required=True, type=Path)
    args = parser.parse_args()

    baseline_rfids = read_baseline_rfids(args.baseline_csv)
    if not baseline_rfids:
        raise SystemExit(f"baseline has no RFID values: {args.baseline_csv}")

    source = load_workbook(args.input_xlsx, read_only=True, data_only=True)
    if args.sheet not in source.sheetnames:
        raise SystemExit(f"sheet {args.sheet!r} not found in {args.input_xlsx}")
    ws = source[args.sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise SystemExit("RFID workbook sheet is empty")

    header = [str(cell or "").strip() for cell in rows[0]]
    try:
        rfid_idx = next(idx for idx, name in enumerate(header) if name.strip().lower() == "rfid")
    except StopIteration as exc:
        raise SystemExit("RFID column not found") from exc

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = args.sheet
    out_ws.append(header)

    excluded_rows: list[dict[str, str]] = []
    included = 0
    considered = 0
    for excel_row_number, row in enumerate(rows[1:], start=2):
        values = list(row)
        rfid = canonical_identifier(values[rfid_idx] if rfid_idx < len(values) else "")
        if not rfid:
            continue
        considered += 1
        if rfid in baseline_rfids:
            out_ws.append(values)
            included += 1
        else:
            excluded_rows.append(
                {
                    "excel_row_number": str(excel_row_number),
                    "rfid": rfid,
                    "reason": "rfid_not_in_trusted_baseline_delta_review",
                }
            )

    args.output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(args.output_xlsx)

    args.excluded_csv.parent.mkdir(parents=True, exist_ok=True)
    with args.excluded_csv.open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["excel_row_number", "rfid", "reason"])
        writer.writeheader()
        writer.writerows(excluded_rows)

    summary = {
        "baseline_rfids": len(baseline_rfids),
        "live_rows_considered": considered,
        "included_rows": included,
        "excluded_delta_review_rows": len(excluded_rows),
    }
    args.summary_json.write_text(json.dumps(summary, indent=2, sort_keys=True) + "\n")
    print(json.dumps(summary, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
