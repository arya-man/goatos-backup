#!/usr/bin/env python3
"""Write source-evidence diagnostics for live replay inputs."""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


NON_ALNUM = re.compile(r"[^A-Z0-9]+")


def canonical_identifier(raw: object | None) -> str:
    value = "" if raw is None else str(raw).strip().replace(",", "").replace(" ", "")
    if not value or value.lower() == "none":
        return ""
    if "e" in value.lower():
        try:
            value = str(int(float(value)))
        except ValueError:
            pass
    elif "." in value:
        left, _, right = value.partition(".")
        if right.strip("0") == "":
            value = left
    return NON_ALNUM.sub("", value.upper())


def normalized_farm(raw: object | None) -> str:
    value = canonical_identifier(raw)
    return {"CJB": "CBE", "BLR": "CPT"}.get(value, value)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="") as f:
        return list(csv.DictReader(f))


def read_rfid_xlsx(path: Path, sheet: str) -> list[dict[str, object]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"missing sheet {sheet!r}; available={wb.sheetnames!r}")
    ws = wb[sheet]
    rows_iter = ws.iter_rows(values_only=True)
    header = ["" if value is None else str(value).strip() for value in next(rows_iter)]
    rows = []
    for row_number, values in enumerate(rows_iter, start=2):
        row = dict(zip(header, values))
        if not any("" if value is None else str(value).strip() for value in row.values()):
            continue
        row["_row_number"] = row_number
        rows.append(row)
    return rows


def build_source_indexes(
    bq_current_csv: Path,
    goats_db_db_csv: Path,
) -> tuple[dict[str, list[dict[str, str]]], dict[tuple[str, str], list[dict[str, str]]]]:
    by_old: dict[str, list[dict[str, str]]] = defaultdict(list)
    by_farm_old: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)

    for source, path, farm_key, goat_key in [
        ("bq_current", bq_current_csv, "farm", "goat_key"),
        ("goats_db_db", goats_db_db_csv, "Farm", "Goat ID"),
    ]:
        for row in read_csv(path):
            old_tag = canonical_identifier(row.get(goat_key))
            farm = normalized_farm(row.get(farm_key))
            if not old_tag:
                continue
            indexed = {**row, "_source": source, "_farm": farm, "_old_tag": old_tag}
            by_old[old_tag].append(indexed)
            by_farm_old[(farm, old_tag)].append(indexed)
    return by_old, by_farm_old


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--rfid-xlsx", required=True, type=Path)
    parser.add_argument("--rfid-sheet", default="Combined")
    parser.add_argument("--baseline-rfid-csv", type=Path)
    parser.add_argument("--bq-current-csv", required=True, type=Path)
    parser.add_argument("--goats-db-db-csv", required=True, type=Path)
    parser.add_argument("--diagnostics-csv-out", required=True, type=Path)
    parser.add_argument("--summary-json-out", required=True, type=Path)
    args = parser.parse_args()

    baseline_rfids: set[str] = set()
    if args.baseline_rfid_csv and args.baseline_rfid_csv.exists():
        for row in read_csv(args.baseline_rfid_csv):
            rfid = canonical_identifier(row.get("RFID"))
            if rfid:
                baseline_rfids.add(rfid)

    by_old, by_farm_old = build_source_indexes(args.bq_current_csv, args.goats_db_db_csv)
    rows = []
    counts: Counter[str] = Counter()
    delta_counts: Counter[str] = Counter()

    for row in read_rfid_xlsx(args.rfid_xlsx, args.rfid_sheet):
        rfid = canonical_identifier(row.get("RFID"))
        if not rfid or rfid == "RFID":
            continue
        old_tag = canonical_identifier(row.get("Old ID"))
        suffix_scope = normalized_farm(row.get("Old ID Suffix"))
        same_farm_sources = sorted({item["_source"] for item in by_farm_old.get((suffix_scope, old_tag), [])})
        any_sources = sorted({item["_source"] for item in by_old.get(old_tag, [])})
        rfid_as_goat_sources = sorted({item["_source"] for item in by_old.get(rfid, [])})

        if same_farm_sources:
            classification = "same_farm_old_tag_evidence"
        elif old_tag and any_sources:
            classification = "old_tag_exists_different_farm"
        elif old_tag:
            classification = "old_tag_not_found_in_bq_or_goats_db"
        else:
            classification = "rfid_only_no_old_tag"
        if rfid_as_goat_sources:
            classification += "+rfid_appears_as_bq_or_goats_db_goat_id"

        delta_state = "not_compared"
        if baseline_rfids:
            delta_state = "existing_in_baseline" if rfid in baseline_rfids else "new_since_baseline"

        counts[classification] += 1
        delta_counts[f"{delta_state}|{classification}"] += 1
        rows.append(
            {
                "row_number": row["_row_number"],
                "delta_state": delta_state,
                "classification": classification,
                "farm": row.get("Farm", ""),
                "old_id": row.get("Old ID", ""),
                "old_id_suffix": row.get("Old ID Suffix", ""),
                "normalized_old_tag": old_tag,
                "normalized_suffix_scope": suffix_scope,
                "rfid": rfid,
                "age": row.get("Age", ""),
                "gender": row.get("Gender", ""),
                "breed": row.get("Breed", ""),
                "shed": row.get("Shed", ""),
                "partition": row.get("Partition", ""),
                "same_farm_sources": "|".join(same_farm_sources),
                "any_old_tag_sources": "|".join(any_sources),
                "rfid_as_goat_sources": "|".join(rfid_as_goat_sources),
            }
        )

    args.diagnostics_csv_out.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "row_number",
        "delta_state",
        "classification",
        "farm",
        "old_id",
        "old_id_suffix",
        "normalized_old_tag",
        "normalized_suffix_scope",
        "rfid",
        "age",
        "gender",
        "breed",
        "shed",
        "partition",
        "same_farm_sources",
        "any_old_tag_sources",
        "rfid_as_goat_sources",
    ]
    with args.diagnostics_csv_out.open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    summary = {
        "rfid_rows": len(rows),
        "baseline_rfid_rows": len(baseline_rfids),
        "by_classification": dict(sorted(counts.items())),
        "by_delta_and_classification": dict(sorted(delta_counts.items())),
    }
    args.summary_json_out.write_text(json.dumps(summary, indent=2, sort_keys=True) + "\n")
    print(json.dumps(summary, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
